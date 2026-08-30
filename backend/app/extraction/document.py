"""
Turns an uploaded module document into plain text for the extraction step.

Handles the formats lecturers actually hand over: PowerPoint decks, Word documents,
PDFs, and plain text. Tables are included — module descriptors very often present
learning outcomes as a table rather than prose, so skipping them would silently drop
most of the content we're looking for.

Text-layer only: a scanned PDF yields nothing here. That's the same OCR gap already
noted in app/ingestion.py.
"""
import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pptx", ".docx", ".pdf", ".txt", ".md"}

# Below this, we almost certainly got an empty text layer rather than a real document.
_MIN_USABLE_CHARS = 40


class DocumentParseError(Exception):
    """Raised when a document can't be turned into usable text."""


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_parts.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        slide_parts.append(" | ".join(cells))
        if slide_parts:
            parts.append(f"--- Slide {index} ---\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Unsupported file type '{suffix or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    try:
        if suffix == ".pptx":
            text = _extract_pptx(data)
        elif suffix == ".docx":
            text = _extract_docx(data)
        elif suffix == ".pdf":
            text = _extract_pdf(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except DocumentParseError:
        raise
    except Exception as exc:  # noqa: BLE001 — surface any parser failure as one clear error
        raise DocumentParseError(f"Could not read '{filename}': {exc}") from exc

    text = text.strip()
    if len(text) < _MIN_USABLE_CHARS:
        raise DocumentParseError(
            f"'{filename}' produced almost no text. If it's a scanned PDF or an "
            "image-only deck, OCR would be needed to read it."
        )
    return text
